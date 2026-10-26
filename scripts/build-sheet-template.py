"""Builds sheet-template/EA-Outreach-Dashboard.xlsx.
Upload it to Google Drive and open it with Google Sheets (File > Save as Google Sheets).
Run: python3 scripts/build-sheet-template.py
"""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

OUT = Path(__file__).resolve().parent.parent / "sheet-template" / "EA-Outreach-Dashboard.xlsx"

HEAD_FILL = PatternFill("solid", fgColor="0F172A")
HEAD_FONT = Font(bold=True, color="FFFFFF")
WRAP_TOP = Alignment(wrap_text=True, vertical="top")

JOBS = [
    ("Job Posting URL", 34), ("Job Title", 30), ("Company Name", 22), ("Location", 22),
    ("Job Description", 60), ("LLM-Generated Outreach Message", 70), ("Status", 16), ("Owner", 14),
    ("Rep Notes", 30), ("Date Posted", 12), ("Date Added", 12), ("Source", 12),
    ("Language Requirements", 28), ("Validation Notes", 40), ("Prompt Version", 14),
]
FILTERED = [
    ("Job Posting URL", 34), ("Job Title", 30), ("Company Name", 22), ("Location", 22), ("Source", 12),
    ("Rejected At", 16), ("Reason", 60), ("Language Requirements", 28), ("Date Checked", 12),
]
PROMPTS = [("Version", 16), ("Active", 10), ("System Prompt", 60), ("User Prompt", 100), ("Notes", 40)]

SYSTEM_V1 = (
    "You write short, specific B2B outreach emails on behalf of a placement company. "
    "Write like a thoughtful person, not a marketer. Never invent facts about the company or the role. "
    "No buzzwords, no exclamation marks, no em dashes. Plain text only."
)
USER_V1 = """Write a cold email from {{sender_name}} at {{sender_company}} to the hiring manager at {{company_name}}, who is hiring a "{{job_title}}" ({{location}}).

What {{sender_company}} offers: {{sender_pitch}}

Details pulled from their posting:
- Who the assistant supports: {{executive_supported}}
- Key responsibilities: {{key_responsibilities}}
- Tools mentioned: {{tools_mentioned}}

Rules:
1. Under 120 words.
2. Open with one sentence that shows you read the posting, using a concrete detail (who they support or a specific responsibility). Do not open with "I hope this finds you well" or "I came across your posting".
3. Explain in one or two sentences how a pre-vetted remote EA from {{sender_company}} covers exactly those responsibilities and tools.
4. One clear ask: a 15 minute call this week to share two or three matching profiles.
5. Sign off as {{sender_name}}, {{sender_company}}.
6. Format: first line "Subject: ..." (under 8 words, specific to their role), then a blank line, then the body.

Full job description for reference:
{{job_description}}"""

SYSTEM_V2 = (
    "You write very short, direct outreach emails. Confident, no fluff, no flattery. "
    "Never invent facts. No exclamation marks, no em dashes. Plain text only."
)
USER_V2 = """Write a 3 sentence cold email from {{sender_name}} at {{sender_company}} to {{company_name}} about their open "{{job_title}}" role.
Sentence 1: name the single most demanding responsibility from their posting ({{key_responsibilities}}).
Sentence 2: {{sender_pitch}}
Sentence 3: ask if they want two matching profiles by Friday.
Format: "Subject: ..." on the first line, blank line, then the body. Sign off as {{sender_name}}.

Job description:
{{job_description}}"""

HOW_TO = [
    "HOW THE SALES REP USES THIS SHEET",
    "",
    "1. New qualified postings land in the Jobs tab every morning (Status = New).",
    "2. Filter the Status column to New. Read the posting (click the URL) and the drafted message.",
    "3. Edit the message if needed, then send it from your own email or LinkedIn InMail to the company's hiring contact.",
    "4. Set Owner to your name and Status to Contacted. Add anything useful in Rep Notes.",
    "5. Update Status as the conversation moves: Replied, Meeting Booked, Not Interested or Skip.",
    "",
    "Never delete rows: the automation uses the URL column to avoid processing the same posting twice.",
    "To hide old rows, filter or sort instead.",
    "",
    "CHANGING THE MESSAGE STYLE (no n8n access needed)",
    "Open the Prompts tab. Set Active to TRUE on the version you want. If two rows are TRUE, each new posting",
    "randomly gets one of them and the Prompt Version column shows which, so you can compare reply rates.",
    "Available placeholders: {{job_title}} {{company_name}} {{location}} {{job_description}} {{executive_supported}}",
    "{{key_responsibilities}} {{tools_mentioned}} {{language_requirements}} {{source}} {{sender_name}}",
    "{{sender_company}} {{sender_pitch}}",
    "",
    "WHY WAS A POSTING SKIPPED?",
    "Check the Filtered Out tab. Every rejected posting is listed with the stage (Rules or LLM validation) and the reason.",
    "",
    "Job data comes from Remotive, Jobicy and Himalayas public APIs. Keep the Source and URL when sharing a listing.",
]


def header(ws, cols):
    for i, (name, width) in enumerate(cols, start=1):
        c = ws.cell(row=1, column=i, value=name)
        c.fill, c.font = HEAD_FILL, HEAD_FONT
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[c.column_letter].width = width
    ws.freeze_panes = "B2" if ws.title == "Jobs" else "A2"
    ws.row_dimensions[1].height = 30


wb = Workbook()
jobs = wb.active
jobs.title = "Jobs"
header(jobs, JOBS)
dv = DataValidation(type="list", formula1='"New,Contacted,Replied,Meeting Booked,Not Interested,Skip"', allow_blank=True)
jobs.add_data_validation(dv)
dv.add("G2:G5000")
for status, color in [("New", "DBEAFE"), ("Contacted", "FEF3C7"), ("Replied", "D1FAE5"),
                      ("Meeting Booked", "A7F3D0"), ("Not Interested", "F3F4F6"), ("Skip", "F3F4F6")]:
    jobs.conditional_formatting.add(
        "A2:O5000", FormulaRule(formula=[f'$G2="{status}"'], fill=PatternFill("solid", fgColor=color)))

filtered = wb.create_sheet("Filtered Out")
header(filtered, FILTERED)

prompts = wb.create_sheet("Prompts")
header(prompts, PROMPTS)
rows = [
    ("v1-consultative", "TRUE", SYSTEM_V1, USER_V1, "Default. Specific, warm, references the posting."),
    ("v2-direct", "FALSE", SYSTEM_V2, USER_V2, "Shorter and blunter. Set Active to TRUE to A/B test against v1."),
]
for r, row in enumerate(rows, start=2):
    for c, val in enumerate(row, start=1):
        cell = prompts.cell(row=r, column=c, value=val)
        cell.alignment = WRAP_TOP
    prompts.row_dimensions[r].height = 260
pdv = DataValidation(type="list", formula1='"TRUE,FALSE"', allow_blank=False)
prompts.add_data_validation(pdv)
pdv.add("B2:B100")

guide = wb.create_sheet("How To Use")
guide.column_dimensions["A"].width = 120
for i, line in enumerate(HOW_TO, start=1):
    c = guide.cell(row=i, column=1, value=line)
    if line.isupper():
        c.font = Font(bold=True, size=12)

OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print("Saved", OUT)
